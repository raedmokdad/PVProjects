from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from filter.facts import facts_from_row
from filter.pv_storage import classify
from portals.click import notice_click_url
from store.db import Store, data_dir


def export_dir() -> Path:
    path = data_dir() / "exports"
    path.mkdir(parents=True, exist_ok=True)
    return path


HEADERS = [
    "Kategorie",
    "Portal",
    "Veröffentlicht",
    "Bewerbungsfrist",
    "Baubeginn",
    "Projektende",
    "Titel",
    "Auftraggeber",
    "Ort",
    "Fläche m²",
    "kWp",
    "Neubau",
    "Trafo",
    "NUTS",
    "Vergabeplattform",
    "Typ",
    "Vergabeordnung",
    "CPV",
    "Treffergrund",
    "Kurztext",
    "Link",
    "pid",
]


def export_excel(store: Store, published_on: str) -> Path:
    out = export_dir()
    rows = store.matches_unique(published_on)
    path_dated = out / f"{published_on}.xlsx"
    path_current = out / "aktuell.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "PV und Speicher"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        facts = facts_from_row(dict(row))
        cat, reason = classify(row["title"] or "", row["excerpt"] or "", row["cpv"] or "")
        # Neubau / Trafo: aus DB-Spalte (gesetzt von KI) oder leer
        is_new_build = row["is_new_build"] if "is_new_build" in row.keys() else None
        has_transformer = row["has_transformer"] if "has_transformer" in row.keys() else None
        neubau_label = ("Ja" if is_new_build else "Nein") if is_new_build is not None else ""
        trafo_label = ("Ja" if has_transformer else "Nein") if has_transformer is not None else ""
        start_on = row["start_on"] if "start_on" in row.keys() else None

        ws.append(
            [
                cat or row["category"],
                row["portal"],
                row["published_on"],
                row["deadline"],
                start_on or "",
                facts["completion_on"] or "",
                row["title"],
                row["organisation"],
                row["city"],
                facts["area_m2"] if facts["area_m2"] is not None else "",
                facts["capacity_kwp"] if facts["capacity_kwp"] is not None else "",
                neubau_label,
                trafo_label,
                row["nuts"] or "",
                row["source_platform"] or "",
                row["notice_type"],
                row["contracting_rule"],
                row["cpv"],
                reason or row["match_reason"],
                row["excerpt"],
                notice_click_url(row["portal"], row["pid"], row["project_url"] or ""),
                row["pid"],
            ]
        )
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    widths = [16, 18, 14, 14, 14, 14, 50, 32, 18, 12, 10, 10, 10, 12, 22, 16, 16, 36, 36, 60, 40, 16]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = width
    for row in ws.iter_rows(min_row=2, max_col=22):
        row[6].alignment = Alignment(wrap_text=True)   # Titel (Spalte G)
        row[19].alignment = Alignment(wrap_text=True)  # Kurztext (Spalte T)

    wb.save(path_dated)
    wb.save(path_current)
    return path_dated
